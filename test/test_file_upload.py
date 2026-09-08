"""Offline regression coverage for attachment selection, upload, and preparation."""

import base64
import copy
import os
import sys
import tempfile
import threading
from pathlib import Path
from unittest.mock import patch

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from claude_chat.clients.base import extract_api_message
from claude_chat.config import MAX_ATTACHMENT_SIZE
from claude_chat.http_router import validate_http_attachment_paths
from claude_chat.services.conversation_service import (
    ConversationService,
    is_attachment_content_block,
    prepare_attachment_content,
)
from claude_chat.services.file_service import FileService
from claude_chat.services.attachment_store import (
    load_managed_attachment,
    store_attachment_path,
)


class MemoryConfig:
    def __init__(self, values=None):
        self.values = values or {}

    def get(self, key, default=None):
        return self.values.get(key, default)


class DummyApp:
    window = None
    config = MemoryConfig()


class MemoryConversationManager:
    def __init__(self, conversation):
        self.conversation = copy.deepcopy(conversation)

    def load_conversation(self, _conv_id):
        return copy.deepcopy(self.conversation)

    def save_conversation(self, conversation):
        self.conversation = copy.deepcopy(conversation)


class EditingApp:
    def __init__(self, conversation):
        self.lock = threading.RLock()
        self.is_streaming = False
        self.conv_manager = MemoryConversationManager(conversation)


def main():
    service = FileService(DummyApp())

    with tempfile.TemporaryDirectory() as temp_dir:
        store_patch = patch(
            "claude_chat.services.attachment_store.ATTACHMENT_STORE_DIR",
            Path(temp_dir) / "attachments",
        )
        store_patch.start()
        try:
            payload = b"hello attachment"
            uploaded = service.upload_dropped_file(
                "../notes.txt",
                len(payload),
                "data:text/plain;base64," + base64.b64encode(payload).decode("ascii"),
            )
            assert not uploaded.get("error")
            assert uploaded["name"] == "notes.txt"
            assert "path" not in uploaded
            stored_path, stored_metadata = load_managed_attachment(uploaded["preview_id"])
            assert stored_path.read_bytes() == payload
            assert stored_path.parent == Path(temp_dir) / "attachments"
            assert stored_metadata["name"] == "notes.txt"

            mismatch = service.upload_dropped_file(
                "bad.txt",
                len(payload) + 1,
                base64.b64encode(payload).decode("ascii"),
            )
            assert "大小不匹配" in mismatch["error"]

            invalid = service.upload_dropped_file("bad.txt", 3, "%%%")
            assert "Base64" in invalid["error"]

            oversized = service.upload_dropped_file("large.txt", MAX_ATTACHMENT_SIZE + 1, "AA==")
            assert "20 MB" in oversized["error"]

            clipboard_file = Path(temp_dir) / "clipboard.txt"
            clipboard_file.write_text("clipboard file", encoding="utf-8")
            with (
                patch("sys.platform", "win32"),
                patch("PIL.ImageGrab.grabclipboard", return_value=[str(clipboard_file)]),
            ):
                clipboard_files = service.paste_attachments_from_clipboard()
            assert clipboard_files["errors"] == []
            assert clipboard_files["attachments"][0]["name"] == "clipboard.txt"

            from PIL import Image

            clipboard_image = Image.new("RGB", (2, 2), color="red")
            with (
                patch("sys.platform", "win32"),
                patch("PIL.ImageGrab.grabclipboard", return_value=clipboard_image),
            ):
                clipboard_images = service.paste_attachments_from_clipboard()
            assert clipboard_images["errors"] == []
            assert clipboard_images["attachments"][0]["name"].endswith(".png")
            image_path, _ = load_managed_attachment(clipboard_images["attachments"][0]["preview_id"])
            assert image_path.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")

        finally:
            store_patch.stop()

        # Start a fresh managed-store patch for attachment preparation coverage.
        store_patch = patch(
            "claude_chat.services.attachment_store.ATTACHMENT_STORE_DIR",
            Path(temp_dir) / "attachments",
        )
        store_patch.start()

        text_path = Path(temp_dir) / "sample.txt"
        text_path.write_text("你好，attachment", encoding="utf-8")
        text_attachment = store_attachment_path(text_path)
        text_block = prepare_attachment_content(
            text_attachment,
            "claude",
            MemoryConfig(),
        )
        assert text_block["type"] == "text"
        assert "你好，attachment" in text_block["text"]
        assert text_block["_attachment"]["name"] == "sample.txt"
        assert text_block["_attachment"]["size"] == text_path.stat().st_size
        assert text_block["_attachment"]["media_type"] == "text/plain"
        assert text_block["_attachment"]["kind"] == "text"
        assert text_block["_attachment"]["preview_id"] == text_attachment["preview_id"]
        assert text_block["_attachment"]["preview_available"] is True
        assert is_attachment_content_block(text_block)
        text_api_message = extract_api_message({"role": "user", "content": [text_block]})
        assert "你好，attachment" in text_api_message["content"][0]["text"]
        assert "_attachment" not in text_api_message["content"][0]

        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation

        docx_path = Path(temp_dir) / "sample.docx"
        document = Document()
        document.add_heading("Word attachment", level=1)
        document.add_paragraph("docx body")
        document.save(docx_path)
        docx_block = prepare_attachment_content(
            store_attachment_path(docx_path),
            "claude",
            MemoryConfig(),
        )
        assert "# Word attachment" in docx_block["text"]
        assert "docx body" in docx_block["text"]
        assert docx_block["_attachment"]["name"] == "sample.docx"
        assert docx_block["_attachment"]["kind"] == "document"

        broken_docx_path = Path(temp_dir) / "private-user-folder" / "broken.docx"
        broken_docx_path.parent.mkdir()
        broken_docx_path.write_bytes(b"not a valid docx")
        broken_docx_block = prepare_attachment_content(
            store_attachment_path(broken_docx_path),
            "claude",
            MemoryConfig(),
        )
        broken_docx_api = extract_api_message({"role": "user", "content": [broken_docx_block]})
        broken_docx_text = broken_docx_api["content"][0]["text"]
        assert "[解析 Word 文档失败]" in broken_docx_text
        assert str(broken_docx_path.parent) not in broken_docx_text
        assert "private-user-folder" not in broken_docx_text

        xlsx_path = Path(temp_dir) / "sample.xlsx"
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.append(["name", "value"])
        worksheet.append(["attachment", 42])
        workbook.save(xlsx_path)
        xlsx_block = prepare_attachment_content(
            store_attachment_path(xlsx_path),
            "gemini",
            MemoryConfig(),
        )
        assert "| name | value |" in xlsx_block["text"]
        assert "| attachment | 42 |" in xlsx_block["text"]

        pptx_path = Path(temp_dir) / "sample.pptx"
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "Slides attachment"
        slide.placeholders[1].text = "pptx body"
        presentation.save(pptx_path)
        pptx_block = prepare_attachment_content(
            store_attachment_path(pptx_path),
            "custom:demo",
            MemoryConfig(),
        )
        assert "Slides attachment" in pptx_block["text"]
        assert "pptx body" in pptx_block["text"]

        image_path = Path(temp_dir) / "pixel.png"
        image_path.write_bytes(b"\x89PNG\r\n\x1a\n")
        image_attachment = store_attachment_path(image_path)
        image_block = prepare_attachment_content(
            image_attachment,
            "claude",
            MemoryConfig(),
        )
        assert image_block["type"] == "image"
        assert image_block["_attachment"]["name"] == "pixel.png"
        assert image_block["_attachment"]["kind"] == "image"
        assert image_block["_attachment"]["preview_id"] == image_attachment["preview_id"]
        api_message = extract_api_message(
            {"role": "user", "content": [image_block]},
            preserve_file_paths=True,
        )
        assert api_message["content"][0]["source"]["file_path"] == str(image_path)
        assert "_attachment" not in api_message["content"][0]

        with patch(
            "claude_chat.attachment_parser.parse_attachment_to_markdown",
            return_value="parsed image",
        ):
            custom_block = prepare_attachment_content(
                image_attachment,
                "custom:demo",
                MemoryConfig(),
            )
        assert custom_block["type"] == "text"
        assert "parsed image" in custom_block["text"]
        assert custom_block["_attachment"]["kind"] == "image"
        custom_api_message = extract_api_message({"role": "user", "content": [custom_block]})
        assert "parsed image" in custom_api_message["content"][0]["text"]
        assert "_attachment" not in custom_api_message["content"][0]

        custom_inline_block = prepare_attachment_content(
            image_attachment,
            "custom:demo",
            MemoryConfig({
                "custom_providers": [{"id": "demo", "provider_adapter": "openrouter"}]
            }),
        )
        assert custom_inline_block["type"] == "image"
        assert custom_inline_block["source"]["file_path"] == str(image_path)

        legacy_block = {
            "type": "text",
            "text": "\n--- 附件文件: legacy.txt ---\nlegacy payload\n--- 附件结束 ---",
        }
        assert is_attachment_content_block(legacy_block)

        editing_app = EditingApp({
            "id": "edit-attachment",
            "messages": [
                {
                    "role": "user",
                    "content": [
                        {"type": "text", "text": "旧问题"},
                        text_block,
                    ],
                },
                {"role": "assistant", "content": "旧回答"},
            ],
        })
        conversation_service = ConversationService(editing_app)
        conversation_service._start_stream_generation = lambda *_args, **_kwargs: True
        assert conversation_service.edit_and_resend("edit-attachment", 0, "新问题")
        edited_messages = editing_app.conv_manager.conversation["messages"]
        assert len(edited_messages) == 1
        assert edited_messages[0]["content"][0] == {"type": "text", "text": "新问题"}
        assert edited_messages[0]["content"][1]["_attachment"]["name"] == "sample.txt"
        assert "你好，attachment" in edited_messages[0]["content"][1]["text"]

        attachment_only_app = EditingApp({
            "id": "edit-attachment-only",
            "messages": [{"role": "user", "content": [text_block]}],
        })
        attachment_only_service = ConversationService(attachment_only_app)
        attachment_only_service._start_stream_generation = lambda *_args, **_kwargs: True
        assert attachment_only_service.edit_and_resend("edit-attachment-only", 0, "")
        attachment_only_content = attachment_only_app.conv_manager.conversation["messages"][0]["content"]
        assert attachment_only_content[0] == {"type": "text", "text": ""}
        assert attachment_only_content[1]["_attachment"]["name"] == "sample.txt"

        unsupported_path = Path(temp_dir) / "archive.zip"
        unsupported_path.write_bytes(b"PK")
        try:
            store_attachment_path(unsupported_path)
            raise AssertionError("unsupported attachment should have been rejected")
        except ValueError as exc:
            assert "不支持" in str(exc)

        # HTTP 边界只接受服务签发的受管附件 ID，任何路径字段都拒绝。
        validate_http_attachment_paths([text_attachment])
        validate_http_attachment_paths([])
        try:
            validate_http_attachment_paths(None)
            raise AssertionError("null attachments should have been rejected")
        except ValueError:
            pass

        secret_path = Path(temp_dir) / "secret.json"
        secret_path.write_text('{"api_key": "leak"}', encoding="utf-8")
        try:
            validate_http_attachment_paths([{"path": str(secret_path)}])
            raise AssertionError("raw paths should have been rejected")
        except ValueError as exc:
            assert "不接受本地路径" in str(exc)

        try:
            validate_http_attachment_paths([{"preview_id": "../../secret"}])
            raise AssertionError("invalid preview IDs should have been rejected")
        except ValueError:
            pass

        try:
            validate_http_attachment_paths("not-a-list")
            raise AssertionError("non-list attachments should have been rejected")
        except ValueError as exc:
            assert "数组" in str(exc)

        store_patch.stop()

    print("file upload regression tests passed")


if __name__ == "__main__":
    main()
